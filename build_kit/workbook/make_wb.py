from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
INK, MIST, INPUT, EX = '12233F', 'EEF3FA', 'FFF4DA', 'F0F0F0'
F = lambda **k: Font(name='Arial', **k)
thin = Side(style='thin', color='D3DDEB'); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def header(ws, row, titles, widths):
    for j, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=j, value=t); c.font = F(bold=True, color='FFFFFF', size=10); c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left'); c.border = BORD; ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[row].height = 42
def title(ws, text, sub):
    ws['A1'] = text; ws['A1'].font = F(bold=True, size=14, color=INK); ws['A2'] = sub; ws['A2'].font = F(size=10, color='55657C'); ws.row_dimensions[2].height = 30
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
def style_inputs(ws, r0, r1, c0, c1, formula_cols=()):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c); cell.font = F(size=10); cell.border = BORD; cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.fill = PatternFill('solid', fgColor=('FFFFFF' if c in formula_cols else INPUT))
def example(ws, row, values):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v); c.font = F(size=10, italic=True, color='55657C'); c.fill = PatternFill('solid', fgColor=EX); c.border = BORD; c.alignment = Alignment(wrap_text=True, vertical='top')
def dv(ws, formula, rng):
    d = DataValidation(type='list', formula1=formula, allow_blank=True); d.error = 'Επιλέξτε τιμή από τη λίστα.'; ws.add_data_validation(d); d.add(rng)

# ---------------- Οδηγίες ----------------
ws = wb.active; ws.title = 'Οδηγίες'; ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 95
ws['A1'] = 'Συστηματική βιβλιογραφική ανασκόπηση: αρχείο προτύπων'; ws['A1'].font = F(bold=True, size=15, color=INK)
rows = [('Σκοπός', 'Ένα φύλλο για κάθε βήμα της ανασκόπησης. Ό,τι καταγράφετε εδώ μεταφέρεται στη Μεθοδολογία και στα Αποτελέσματα της εργασίας.'),
 ('Πού γράφετε', 'Μόνο στα κελιά με κίτρινο φόντο. Τα λευκά κελιά περιέχουν τύπους και υπολογίζονται αυτόματα.'),
 ('Γραμμές παραδείγματος', 'Οι γκρίζες γραμμές με πλάγια γράμματα είναι παραδείγματα. Διαγράψτε το περιεχόμενό τους πριν ξεκινήσετε: αλλιώς προσμετρώνται στους αριθμούς του PRISMA.'),
 ('1_Αναζήτηση', 'Ημερολόγιο αναζήτησης. Μία γραμμή για κάθε αναζήτηση σε κάθε βάση δεδομένων, με τη συμβολοσειρά ακριβώς όπως εκτελέστηκε.'),
 ('2_Επιλογή', 'Μία γραμμή για κάθε εγγραφή που εντοπίστηκε. Τα δύο μέλη αποφασίζουν ανεξάρτητα στο Στάδιο 1 (τίτλος και περίληψη). Η στήλη «Συμφωνία» δείχνει πού χρειάζεται συζήτηση.'),
 ('PRISMA_αριθμοί', 'Υπολογίζει αυτόματα τους αριθμούς του διαγράμματος ροής PRISMA 2020 από το φύλλο 2_Επιλογή και ελέγχει ότι αθροίζονται.'),
 ('3_Εξαγωγή', 'Πίνακας εξαγωγής δεδομένων. Μία γραμμή ανά μελέτη που εντάχθηκε. Οι στήλες ακολουθούν τις τέσσερις ερωτήσεις του μαθήματος: Τι, Ποιος, Γιατί, Πώς.'),
 ('4_Ποιότητα', 'Πέντε ερωτήσεις αξιολόγησης ανά μελέτη, με απάντηση ΝΑΙ, ΕΝ ΜΕΡΕΙ ή ΟΧΙ. Ο βαθμός (0 έως 5) υπολογίζεται αυτόματα: ΝΑΙ = 1, ΕΝ ΜΕΡΕΙ = 0,5, ΟΧΙ = 0.'),
 ('Πηγές', 'Διάγραμμα ροής και κατάλογος ελέγχου: Page et al. (2021), The PRISMA 2020 statement, BMJ. Για πληρέστερη αξιολόγηση ποιότητας: Hong et al. (2018), Mixed Methods Appraisal Tool (MMAT).')]
for i, (a, b) in enumerate(rows, 3):
    ws.cell(row=i, column=1, value=a).font = F(bold=True, size=11, color=INK); c = ws.cell(row=i, column=2, value=b); c.font = F(size=11); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=1).alignment = Alignment(vertical='top'); ws.row_dimensions[i].height = 48

# ---------------- 1_Αναζήτηση ----------------
s1 = wb.create_sheet('1_Αναζήτηση'); title(s1, 'Ημερολόγιο αναζήτησης', 'Μία γραμμή ανά αναζήτηση. Αντιγράψτε τη συμβολοσειρά όπως ακριβώς εκτελέστηκε στη βάση.')
header(s1, 4, ['Α/Α', 'Βάση δεδομένων', 'Ημερομηνία αναζήτησης', 'Συμβολοσειρά αναζήτησης (όπως εκτελέστηκε)', 'Φίλτρα', 'Αποτελέσματα (n)', 'Σημειώσεις'], [6, 20, 16, 70, 30, 16, 36])
example(s1, 5, [1, 'Scopus', '2026-12-10', 'TITLE-ABS-KEY(("learning analytics") AND (dashboard* OR "visual analytics") AND ("self-regulat*" OR SRL) AND ("higher education" OR universit* OR undergraduate*)) AND PUBYEAR > 2018 AND LANGUAGE(english)', 'έτος 2019 και μετά, αγγλικά, άρθρα και συνέδρια', 230, 'Παράδειγμα: διαγράψτε το πριν ξεκινήσετε.'])
style_inputs(s1, 6, 25, 1, 7); s1.freeze_panes = 'A5'
s1['E27'] = 'Σύνολο εγγραφών που εντοπίστηκαν'; s1['E27'].font = F(bold=True, size=10); s1['F27'] = '=SUM(F5:F25)'; s1['F27'].font = F(bold=True, size=11); s1['F27'].border = BORD
dv(s1, '"Scopus,ERIC,IEEE Xplore,ACM Digital Library,Web of Science,Άλλη"', 'B6:B25')

# ---------------- 2_Επιλογή ----------------
s2 = wb.create_sheet('2_Επιλογή'); title(s2, 'Επιλογή μελετών σε δύο στάδια', 'Στάδιο 1: τίτλος και περίληψη, ανεξάρτητα από τα δύο μέλη. Στάδιο 2: πλήρες κείμενο. Κάθε αποκλεισμός στο Στάδιο 2 χρειάζεται λόγο.')
header(s2, 4, ['ID', 'Συγγραφείς', 'Έτος', 'Τίτλος', 'Βάση δεδομένων', 'Διπλοεγγραφή;', 'Στάδιο 1: Μέλος Α', 'Στάδιο 1: Μέλος Β', 'Συμφωνία;', 'Στάδιο 1: τελική απόφαση', 'Πλήρες κείμενο βρέθηκε;', 'Στάδιο 2: τελική απόφαση', 'Λόγος αποκλεισμού (Στάδιο 2)'], [6, 26, 8, 60, 16, 14, 14, 14, 12, 16, 16, 18, 34])
ex = [[1, 'Herodotou, C., Rienties, B., Boroowa, A., Zdrahal, Z., & Hlosta, M.', 2019, 'A large-scale implementation of predictive learning analytics in higher education: the teachers\' role and perspective', 'Scopus', 'ΟΧΙ', 'ΝΑΙ', 'ΝΑΙ', None, 'ΝΑΙ', 'ΝΑΙ', 'ΕΝΤΑΞΗ', ''],
      [2, 'Herodotou, C., Rienties, B., Boroowa, A., Zdrahal, Z., & Hlosta, M.', 2019, 'A large-scale implementation of predictive learning analytics in higher education: the teachers\' role and perspective', 'ERIC', 'ΝΑΙ', '', '', None, '', '', '', ''],
      [3, 'Kaliisa, R., Misiejuk, K., López-Pernas, S., Khalil, M., & Saqr, M.', 2024, 'Have learning analytics dashboards lived up to the hype? A systematic review of impact on students\' achievement, motivation, participation and attitude', 'ACM Digital Library', 'ΟΧΙ', 'ΝΑΙ', 'ΟΧΙ', None, 'ΝΑΙ', 'ΝΑΙ', 'ΑΠΟΚΛΕΙΣΜΟΣ', 'Ανασκόπηση']]
FIRST, LAST = 5, 404
for k, r in enumerate(ex): example(s2, FIRST + k, r)
style_inputs(s2, FIRST + 3, LAST, 1, 13, formula_cols=(9,))
for r in range(FIRST, LAST + 1):
    s2.cell(row=r, column=9, value=f'=IF(OR(G{r}="",H{r}=""),"",IF(G{r}=H{r},"ΝΑΙ","ΣΥΖΗΤΗΣΗ"))')
    if r < FIRST + 3: s2.cell(row=r, column=9).font = F(size=10, italic=True, color='55657C')
s2.freeze_panes = 'E5'
dv(s2, '"ΝΑΙ,ΟΧΙ"', f'F{FIRST}:F{LAST}'); dv(s2, '"ΝΑΙ,ΟΧΙ,ΙΣΩΣ"', f'G{FIRST}:H{LAST}'); dv(s2, '"ΝΑΙ,ΟΧΙ"', f'J{FIRST}:K{LAST}'); dv(s2, '"ΕΝΤΑΞΗ,ΑΠΟΚΛΕΙΣΜΟΣ"', f'L{FIRST}:L{LAST}')
dv(s2, '"Όχι εμπειρική μελέτη,Εκτός πληθυσμού ή πλαισίου,Εκτός έννοιας ενδιαφέροντος,Ανασκόπηση, όχι εμπειρική μελέτη"'.replace('Ανασκόπηση, όχι εμπειρική μελέτη', 'Ανασκόπηση'), f'M{FIRST + 3}:M{LAST}')

# ---------------- PRISMA_αριθμοί ----------------
sp = wb.create_sheet('PRISMA_αριθμοί'); title(sp, 'Αριθμοί για το διάγραμμα ροής PRISMA 2020', 'Υπολογίζονται αυτόματα από το φύλλο 2_Επιλογή. Μην γράφετε σε αυτό το φύλλο.')
sp.column_dimensions['A'].width = 16; sp.column_dimensions['B'].width = 62; sp.column_dimensions['C'].width = 12; sp.column_dimensions['D'].width = 50
R = f"'2_Επιλογή'!"
lines = [('Εντοπισμός', 'Εγγραφές που εντοπίστηκαν (γραμμές στο φύλλο 2_Επιλογή)', f'=COUNTA({R}B{FIRST}:B{LAST})'),
 ('', 'Αφαιρέθηκαν πριν από τον έλεγχο: διπλοεγγραφές', f'=COUNTIF({R}F{FIRST}:F{LAST},"ΝΑΙ")'),
 ('Έλεγχος', 'Εγγραφές που ελέγχθηκαν με βάση τίτλο και περίληψη', '=C5-C6'),
 ('', 'Αποκλείστηκαν στο Στάδιο 1', f'=COUNTIF({R}J{FIRST}:J{LAST},"ΟΧΙ")'),
 ('', 'Πλήρη κείμενα που αναζητήθηκαν', f'=COUNTIF({R}J{FIRST}:J{LAST},"ΝΑΙ")'),
 ('', 'Πλήρη κείμενα που δεν ανακτήθηκαν', f'=COUNTIF({R}K{FIRST}:K{LAST},"ΟΧΙ")'),
 ('', 'Πλήρη κείμενα που αξιολογήθηκαν', '=C9-C10'),
 ('', 'Αποκλείστηκαν με αιτιολογία στο Στάδιο 2', f'=COUNTIF({R}L{FIRST}:L{LAST},"ΑΠΟΚΛΕΙΣΜΟΣ")'),
 ('Ένταξη', 'Μελέτες που εντάχθηκαν στην ανασκόπηση', f'=COUNTIF({R}L{FIRST}:L{LAST},"ΕΝΤΑΞΗ")')]
header(sp, 4, ['Φάση', 'Μέγεθος', 'n', 'Έλεγχος'], [16, 62, 12, 50])
for i, (a, b, f) in enumerate(lines, 5):
    sp.cell(row=i, column=1, value=a).font = F(bold=True, size=11, color=INK); sp.cell(row=i, column=2, value=b).font = F(size=11); c = sp.cell(row=i, column=3, value=f); c.font = F(bold=True, size=12)
    for j in (1, 2, 3, 4): sp.cell(row=i, column=j).border = BORD
sp['D5'] = "=IF(C5='1_Αναζήτηση'!F27,\"Συμφωνεί με το ημερολόγιο αναζήτησης\",\"Διαφέρει από το ημερολόγιο: ελέγξτε ότι καταχωρίσατε όλες τις εγγραφές\")"
sp['D7'] = '=IF(C7=C8+C9,"Αθροίζει","Εκκρεμούν αποφάσεις στο Στάδιο 1")'
sp['D11'] = '=IF(C11=C12+C13,"Αθροίζει","Εκκρεμούν αποφάσεις στο Στάδιο 2")'
for a in ('D5', 'D7', 'D11'): sp[a].font = F(size=10, italic=True)
sp['A16'] = 'Λόγοι αποκλεισμού στο Στάδιο 2'; sp['A16'].font = F(bold=True, size=12, color=INK)
for i, reason in enumerate(['Όχι εμπειρική μελέτη', 'Εκτός πληθυσμού ή πλαισίου', 'Εκτός έννοιας ενδιαφέροντος', 'Ανασκόπηση'], 17):
    sp.cell(row=i, column=2, value=reason).font = F(size=11); sp.cell(row=i, column=3, value=f'=COUNTIF({R}M{FIRST}:M{LAST},B{i})').font = F(bold=True, size=11)
sp['A23'] = 'Μελέτες που εντάχθηκαν, ανά έτος (από το φύλλο 3_Εξαγωγή)'; sp['A23'].font = F(bold=True, size=12, color=INK)
for i, y in enumerate(range(2019, 2027), 24):
    sp.cell(row=i, column=2, value=y).font = F(size=11); sp.cell(row=i, column=3, value=f"=COUNTIF('3_Εξαγωγή'!C5:C104,B{i})").font = F(bold=True, size=11)
sp['D24'] = 'Υπόθεση: χρονικό εύρος 2019 έως 2026, όπως στο παράδειγμα του μαθήματος. Αλλάξτε τα έτη στη στήλη Β αν το δικό σας εύρος διαφέρει.'; sp['D24'].font = F(size=10, italic=True); sp['D24'].alignment = Alignment(wrap_text=True, vertical='top')
for i in range(24, 32): sp.cell(row=i, column=2).fill = PatternFill('solid', fgColor=INPUT)

# ---------------- 3_Εξαγωγή ----------------
s3 = wb.create_sheet('3_Εξαγωγή'); title(s3, 'Πίνακας εξαγωγής δεδομένων', 'Μία γραμμή ανά μελέτη που εντάχθηκε. Ίδιες στήλες για όλες τις μελέτες. Δοκιμάστε τον πίνακα σε τρεις μελέτες πριν συνεχίσετε.')
header(s3, 4, ['ID', 'Παραπομπή (APA 7)', 'Έτος', 'Χώρα', 'Τι; Εκπαιδευτικό πλαίσιο', 'Τι; Δεδομένα', 'Ποιος; Ενδιαφερόμενοι', 'Γιατί; Στόχος', 'Πώς; Μέθοδος', 'Δείγμα', 'Κύρια ευρήματα', 'Περιορισμοί', 'Σχετικό με ΕΕ1;', 'Σχετικό με ΕΕ2;', 'Σχετικό με ΕΕ3;'], [6, 46, 8, 14, 28, 26, 22, 28, 28, 22, 44, 30, 12, 12, 12])
example(s3, 5, [1, 'Herodotou, C., Rienties, B., Boroowa, A., Zdrahal, Z., & Hlosta, M. (2019). A large-scale implementation of predictive learning analytics in higher education: the teachers\' role and perspective. Educational Technology Research and Development, 67(5), 1273–1306.', 2019, 'Ηνωμένο Βασίλειο', 'εξ αποστάσεως πανεπιστήμιο', 'δημογραφικά και ίχνη από το εικονικό περιβάλλον μάθησης', 'διδάσκοντες', 'έγκαιρος εντοπισμός φοιτητών σε κίνδυνο', 'προγνωστικό μοντέλο και πίνακας πληροφοριών', 'διδάσκοντες και οι φοιτητές τους', 'Η συχνότερη χρήση των προβλέψεων από τους διδάσκοντες συνδέθηκε με καλύτερα ποσοστά ολοκλήρωσης και επιτυχίας.', 'Συσχέτιση, όχι αιτιότητα.', 'ΝΑΙ', 'ΝΑΙ', 'ΝΑΙ'])
style_inputs(s3, 6, 104, 1, 15); s3.freeze_panes = 'C5'; dv(s3, '"ΝΑΙ,ΟΧΙ"', 'M6:O104')

# ---------------- 4_Ποιότητα ----------------
s4 = wb.create_sheet('4_Ποιότητα'); title(s4, 'Αξιολόγηση ποιότητας', 'Απαντήστε ΝΑΙ, ΕΝ ΜΕΡΕΙ ή ΟΧΙ. Ο βαθμός υπολογίζεται αυτόματα: ΝΑΙ = 1, ΕΝ ΜΕΡΕΙ = 0,5, ΟΧΙ = 0. Η αξιολόγηση δεν αποκλείει μελέτες: δείχνει πόσο βάρος δίνετε στα ευρήματά τους.')
header(s4, 4, ['ID', 'Παραπομπή (σύντομη)', '1. Σαφές ερευνητικό ερώτημα;', '2. Περιγράφονται δείγμα και πλαίσιο;', '3. Κατάλληλη μέθοδος;', '4. Τα συμπεράσματα στηρίζονται στα δεδομένα;', '5. Αναφέρονται περιορισμοί;', 'Βαθμός (0 έως 5)', 'Σχόλια'], [6, 34, 18, 18, 16, 22, 18, 14, 44])
example(s4, 5, [1, 'Herodotou et al. (2019)', 'ΝΑΙ', 'ΝΑΙ', 'ΝΑΙ', 'ΕΝ ΜΕΡΕΙ', 'ΝΑΙ', None, 'Παράδειγμα: διαγράψτε το πριν ξεκινήσετε.'])
style_inputs(s4, 6, 104, 1, 9, formula_cols=(8,))
for r in range(5, 105):
    c = s4.cell(row=r, column=8, value=f'=IF(COUNTA(C{r}:G{r})=0,"",COUNTIF(C{r}:G{r},"ΝΑΙ")+0.5*COUNTIF(C{r}:G{r},"ΕΝ ΜΕΡΕΙ"))'); c.font = F(bold=True, size=10, italic=(r == 5), color=('55657C' if r == 5 else '000000')); c.border = BORD
s4.freeze_panes = 'C5'; dv(s4, '"ΝΑΙ,ΕΝ ΜΕΡΕΙ,ΟΧΙ"', 'C6:G104')

wb.save('/home/claude/wb/LA_review_templates.xlsx'); print('saved')
