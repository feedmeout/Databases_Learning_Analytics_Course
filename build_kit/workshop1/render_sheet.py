"""Renders the image of sheet 0_Συμβολοσειρά used on slide 18 (img_w2p3/sheet0.png): fills the Week 1 example into the yellow cells of a temporary copy, prints it to PDF with LibreOffice and crops the page. Needs LibreOffice, poppler (pdftoppm) and the xlsx skill's helper scripts."""
import subprocess, sys, glob, os
from openpyxl import load_workbook
from PIL import Image, ImageChops
HERE = os.path.dirname(os.path.abspath(__file__)); K = os.path.join(HERE, '..'); SK = '/mnt/skills/public/xlsx/scripts'
w = load_workbook(f'{K}/out/LA_review_templates_v2.xlsx')
for n in w.sheetnames:
    if n != '0_Συμβολοσειρά': del w[n]
t = w['0_Συμβολοσειρά']
EX = [['learning analytics', 'learning analytics'], ['πίνακες πληροφοριών', 'dashboard*', 'visual analytics'], ['αυτορρύθμιση', 'self-regulated learning', 'self-regulation', 'SRL'], ['ανώτατη εκπαίδευση', 'higher education', 'universit*', 'undergraduate*']]
for j, col in enumerate(EX, start=2):
    for k, v in enumerate(col): t.cell(row=5 + k, column=j, value=v)
t['B16'] = 2019; t['B17'] = 'english'
for r in list(range(10, 14)) + list(range(21, 48)): t.row_dimensions[r].hidden = True          # unused term rows, generic form, checks, built-in example
t.column_dimensions['F'].hidden = True; t.column_dimensions['A'].width = 30
for r in (19, 20):
    t.unmerge_cells(start_row=r, start_column=2, end_row=r, end_column=6); t.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); t.row_dimensions[r].height = 44
t.row_dimensions[14].height = 34
t.print_area = 'A4:E20'; t.page_setup.orientation = 'landscape'; t.page_setup.fitToWidth = 1; t.page_setup.fitToHeight = 1; t.sheet_properties.pageSetUpPr.fitToPage = True
t.page_margins.left = t.page_margins.right = 0.2; t.page_margins.top = t.page_margins.bottom = 0.2
w.save('/tmp/sheet0_shot.xlsx')
subprocess.run(['python3', f'{SK}/recalc.py', '/tmp/sheet0_shot.xlsx', '90'], capture_output=True)
sys.path.insert(0, SK)
from office.soffice import run_soffice
run_soffice(['--headless', '--convert-to', 'pdf', '--outdir', '/tmp', '/tmp/sheet0_shot.xlsx'])
for f in glob.glob('/tmp/sheet0-*.png'): os.remove(f)
subprocess.run(['pdftoppm', '-png', '-r', '240', '/tmp/sheet0_shot.pdf', '/tmp/sheet0'])
im = Image.open(sorted(glob.glob('/tmp/sheet0-*.png'))[0]).convert('RGB'); b = ImageChops.difference(im, Image.new('RGB', im.size, '#fff')).getbbox()
im = im.crop((max(0, b[0] - 12), max(0, b[1] - 12), min(im.size[0], b[2] + 12), min(im.size[1], b[3] + 12)))
os.makedirs(f'{K}/img_w2p3', exist_ok=True); im.save(f'{K}/img_w2p3/sheet0.png'); print('sheet0.png', im.size)
